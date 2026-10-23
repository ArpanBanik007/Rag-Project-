
import os
import json
import time
from pathlib import Path
from typing import Optional

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# ============================================================
# 1. CONFIGURATION
# ============================================================

load_dotenv()

API_KEY = os.getenv("GROQ_API_KEY")

if not API_KEY:
    raise ValueError(
        "GROQ_API_KEY not found. Please check your .env file."
    )

client = Groq(api_key=API_KEY)

MODEL = "openai/gpt-oss-20b"

BASE_DIR = Path(__file__).resolve().parent
RESUME_FOLDER = BASE_DIR / "resumes"
OUTPUT_FOLDER = BASE_DIR / "output"

OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

JOB_JSON_PATH = OUTPUT_FOLDER / "job_description.json"
RESULTS_JSON_PATH = OUTPUT_FOLDER / "candidate_results.json"

# Limits the amount of text sent to the model.
# Increase cautiously if your resumes are unusually long.
MAX_RESUME_CHARS = 18000
MAX_JD_CHARS = 10000


# ============================================================
# 2. PYDANTIC SCHEMAS
# ============================================================

class Experience(BaseModel):
    company: Optional[str] = None
    position: Optional[str] = None
    duration: Optional[str] = None
    description: str = ""


class Education(BaseModel):
    degree: Optional[str] = None
    institution: Optional[str] = None
    graduation_year: Optional[str] = None


class Project(BaseModel):
    name: Optional[str] = None
    description: str = ""
    technologies: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    name: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None

    total_experience_years: float = 0.0

    skills: list[str] = Field(default_factory=list)
    experiences: list[Experience] = Field(default_factory=list)
    education: list[Education] = Field(default_factory=list)
    projects: list[Project] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


class JobD(BaseModel):
    role: str
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    minimum_experience: float = 0.0
    education_requirements: list[str] = Field(default_factory=list)
    responsibilities: list[str] = Field(default_factory=list)


# ============================================================
# 3. GROQ JSON HELPER
# ============================================================

def ask_groq_json(
    system_prompt: str,
    user_prompt: str,
    max_tokens: int = 8192
) -> dict:
    """
    Ask Groq for a JSON object and validate the response.

    Raises an exception if the response is empty or invalid.
    """

    response = client.chat.completions.create(
        model=MODEL,
        messages=[
            {
                "role": "system",
                "content": system_prompt
            },
            {
                "role": "user",
                "content": user_prompt
            }
        ],
        response_format={"type": "json_object"},
        max_completion_tokens=max_tokens,
        temperature=0.1
    )

    message = response.choices[0].message
    raw_output = message.content

    if not raw_output:
        raise ValueError("Groq returned an empty response.")

    try:
        return json.loads(raw_output)
    except json.JSONDecodeError as exc:
        finish_reason = response.choices[0].finish_reason

        raise ValueError(
            f"Groq returned invalid JSON. "
            f"Finish reason: {finish_reason}. "
            f"Output length: {len(raw_output)} characters."
        ) from exc


# ============================================================
# 4. READ PDF AND DOCX FILES
# ============================================================

def read_resume(file_path: Path) -> str:
    """Extract text from a PDF or DOCX resume."""

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        reader = PdfReader(str(file_path))

        pages = []

        for page in reader.pages:
            pages.append(page.extract_text() or "")

        text = "\n".join(pages)

    elif extension == ".docx":
        document = Document(str(file_path))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        tables = []

        for table in document.tables:
            for row in table.rows:
                tables.append(
                    " | ".join(cell.text.strip() for cell in row.cells)
                )

        text = "\n".join(paragraphs + tables)

    else:
        raise ValueError(
            f"Unsupported resume format: {extension}"
        )

    if not text.strip():
        raise ValueError(
            "No readable text found. The PDF may be scanned "
            "or image-based and may require OCR."
        )

    return text.strip()


# ============================================================
# 5. PARSE JOB DESCRIPTION
# ============================================================

def parse_job_description(job_description: str) -> JobD:
    """Extract structured requirements from a job description."""

    job_description = job_description[:MAX_JD_CHARS]

    system_prompt = """
You are a job-description information extraction system.

Extract the job requirements into one valid JSON object.

Use exactly these fields:
{
  "role": "string",
  "required_skills": ["string"],
  "preferred_skills": ["string"],
  "minimum_experience": 0,
  "education_requirements": ["string"],
  "responsibilities": ["string"]
}

Rules:
- Return JSON only.
- Do not use Markdown or code fences.
- Do not invent requirements.
- Use an empty list if information is missing.
- Use 0 for minimum_experience if it is unspecified.
- Keep responsibilities concise.
- Use numeric years for minimum_experience.
"""

    data = ask_groq_json(
        system_prompt=system_prompt,
        user_prompt=(
            "Extract the information from this job description:\n\n"
            + job_description
        ),
        max_tokens=4096
    )

    return JobD(**data)


# ============================================================
# 6. PARSE RESUME
# ============================================================

def parse_resume(resume_text: str) -> Resume:
    """Extract structured candidate information from resume text."""

    # Limit the input so unusually long resumes do not consume
    # the entire model context.
    resume_text = resume_text[:MAX_RESUME_CHARS]

    system_prompt = """
You are an accurate resume-parsing system.

Extract information from the supplied resume into one valid JSON object.

Use exactly these fields:
{
  "name": null,
  "email": null,
  "phone": null,
  "total_experience_years": 0,
  "skills": [],
  "experiences": [
    {
      "company": null,
      "position": null,
      "duration": null,
      "description": ""
    }
  ],
  "education": [
    {
      "degree": null,
      "institution": null,
      "graduation_year": null
    }
  ],
  "projects": [
    {
      "name": null,
      "description": "",
      "technologies": []
    }
  ],
  "certifications": []
}

Rules:
- Return JSON only, without Markdown or code fences.
- Do not invent or assume candidate information.
- Use null for missing single-value fields.
- Use [] for missing lists.
- Use 0 for unknown total_experience_years.
- Keep experience descriptions concise.
- Keep project descriptions concise.
- Include relevant skills only.
- Do not duplicate skills.
- Experience years must be a number.
- Do not treat academic projects as employment.
- Preserve the actual information in the resume.
- Avoid unnecessary explanations.
"""

    data = ask_groq_json(
        system_prompt=system_prompt,
        user_prompt=(
            "Parse the following resume into the specified JSON structure.\n\n"
            "RESUME TEXT:\n"
            + resume_text
        ),
        max_tokens=8192
    )

    return Resume(**data)


# ============================================================
# 7. NORMALIZE SKILLS
# ============================================================

def normalize_skill(skill: str) -> str:
    """Normalize skill text for basic case-insensitive matching."""

    return " ".join(skill.lower().strip().split())


def get_matched_skills(
    required_skills: list[str],
    candidate_skills: list[str]
) -> list[str]:
    """
    Match required skills against extracted candidate skills.

    Matching is case-insensitive. This is a simple baseline and
    does not identify all semantic equivalents.
    """

    candidate_normalized = {
        normalize_skill(skill)
        for skill in candidate_skills
        if skill.strip()
    }

    matched = []

    for skill in required_skills:
        normalized = normalize_skill(skill)

        if normalized in candidate_normalized:
            matched.append(skill)

    return matched


# ============================================================
# 8. SCORE CANDIDATE
# ============================================================

def final_score(job: JobD, resume: Resume) -> dict:
    """
    Calculate a transparent baseline score.

    Weights:
    - Required skills: 60 points
    - Experience: 25 points
    - Education information: 15 points

    This is a screening aid, not a hiring decision.
    """

    matched_required = get_matched_skills(
        job.required_skills,
        resume.skills
    )

    missing_required = [
        skill
        for skill in job.required_skills
        if skill not in matched_required
    ]

    matched_preferred = get_matched_skills(
        job.preferred_skills,
        resume.skills
    )

    # Required skills: 60 points.
    if job.required_skills:
        skill_score = (
            len(matched_required) / len(job.required_skills)
        ) * 60
    else:
        skill_score = 60.0

    # Experience: 25 points.
    if job.minimum_experience > 0:
        experience_score = min(
            resume.total_experience_years
            / job.minimum_experience,
            1.0
        ) * 25
    else:
        experience_score = 25.0

    # Education: 15 points.
    # This baseline checks whether education information exists.
    # It does not claim that the candidate's degree satisfies
    # every job-specific education requirement.
    if not job.education_requirements:
        education_score = 15.0
        education_status = "No specific education requirement provided"
    elif resume.education:
        education_score = 15.0
        education_status = "Education information available; verify requirements"
    else:
        education_score = 0.0
        education_status = "Education information not found"

    total_score = round(
        skill_score + experience_score + education_score,
        2
    )

    return {
        "score": total_score,
        "score_breakdown": {
            "required_skills": round(skill_score, 2),
            "experience": round(experience_score, 2),
            "education_information": round(education_score, 2)
        },
        "matched_required_skills": matched_required,
        "missing_required_skills": missing_required,
        "matched_preferred_skills": matched_preferred,
        "candidate_experience_years": resume.total_experience_years,
        "minimum_required_experience_years": job.minimum_experience,
        "education_status": education_status
    }


# ============================================================
# 9. SAVE JSON SAFELY
# ============================================================

def save_json(file_path: Path, data: dict | list) -> None:
    """Save JSON output using UTF-8 encoding."""

    with file_path.open("w", encoding="utf-8") as file:
        json.dump(
            data,
            file,
            indent=4,
            ensure_ascii=False
        )


# ============================================================
# 10. MAIN PROGRAM
# ============================================================

def main() -> None:
    print("=" * 60)
    print("RESUME PARSER + JOB DESCRIPTION MATCHER")
    print("=" * 60)

    if not RESUME_FOLDER.exists():
        print(f"Resume folder not found: {RESUME_FOLDER}")
        print("Create a folder named 'resumes' beside resume.py.")
        return

    # --------------------------------------------------------
    # Enter or replace this with the actual job description.
    # --------------------------------------------------------

    job_description = """
    Job Title: Software Development Engineer I (SDE-I)

    Company: Amazon

    Job Description:
    Develop, test, and maintain software applications.
    Write readable and maintainable code.
    Work with team members to design and implement solutions.
    Debug software problems and improve application performance.

    Required Skills:
    Programming fundamentals, data structures, algorithms,
    problem solving, object-oriented programming.

    Preferred Skills:
    Python, Java, C++, SQL, databases, web development.

    Experience:
    Entry-level role. Relevant academic projects and internships
    may be useful. Verify the actual job posting for exact criteria.

    Education:
    Bachelor's degree in Computer Science, Engineering,
    or a related field, subject to the actual job requirements.
    """

    # --------------------------------------------------------
    # Parse job description
    # --------------------------------------------------------

    print("\nParsing job description...")

    try:
        job = parse_job_description(job_description)
        save_json(JOB_JSON_PATH, job.model_dump())

        print("Job description parsed successfully.")
        print("Role:", job.role)
        print("Required skills:", job.required_skills)

    except Exception as error:
        print("Failed to parse job description.")
        print(f"Error: {error}")
        return

    # --------------------------------------------------------
    # Find resumes
    # --------------------------------------------------------

    resume_files = sorted(
        file_path
        for file_path in RESUME_FOLDER.iterdir()
        if file_path.is_file()
        and file_path.suffix.lower() in {".pdf", ".docx"}
    )

    if not resume_files:
        print("\nNo PDF or DOCX resumes found in:")
        print(RESUME_FOLDER)
        return

    print(f"\nFound {len(resume_files)} resume(s).")

    candidate_results = []
    failed_resumes = []

    # --------------------------------------------------------
    # Process resumes one by one
    # --------------------------------------------------------

    for index, file_path in enumerate(resume_files, start=1):
        print("\n" + "-" * 60)
        print(f"[{index}/{len(resume_files)}] Processing: {file_path.name}")

        try:
            # Step A: Extract text.
            resume_text = read_resume(file_path)

            print(f"Extracted {len(resume_text)} characters.")

            # Step B: Parse resume.
            print("Parsing resume with Groq...")

            parsed_resume = parse_resume(resume_text)

            print("Resume parsed successfully.")

            # Step C: Score candidate.
            score_result = final_score(job, parsed_resume)

            candidate_record = {
                "file_name": file_path.name,
                "candidate": parsed_resume.model_dump(),
                "matching_result": score_result
            }

            candidate_results.append(candidate_record)

            print("Candidate:", parsed_resume.name or "Unknown")
            print("Score:", score_result["score"], "/ 100")
            print(
                "Matched required skills:",
                score_result["matched_required_skills"]
            )
            print(
                "Missing required skills:",
                score_result["missing_required_skills"]
            )

        except Exception as error:
            error_message = str(error)

            print(f"Failed to process {file_path.name}")
            print(f"Reason: {error_message}")

            failed_resumes.append({
                "file_name": file_path.name,
                "error": error_message
            })

            # Keep processing the remaining resumes.
            continue

        # Small delay between successful resumes.
        # Increase this if you encounter rate limits.
        if index < len(resume_files):
            time.sleep(2)

    # --------------------------------------------------------
    # Rank candidates
    # --------------------------------------------------------

    candidate_results.sort(
        key=lambda item: item["matching_result"]["score"],
        reverse=True
    )

    for rank, candidate in enumerate(candidate_results, start=1):
        candidate["rank"] = rank

    # --------------------------------------------------------
    # Save results
    # --------------------------------------------------------

    output_data = {
        "job_role": job.role,
        "total_resumes_found": len(resume_files),
        "successfully_processed": len(candidate_results),
        "failed_count": len(failed_resumes),
        "candidates": candidate_results,
        "failed_resumes": failed_resumes
    }

    save_json(RESULTS_JSON_PATH, output_data)

    # --------------------------------------------------------
    # Final report
    # --------------------------------------------------------

    print("\n" + "=" * 60)
    print("PROCESSING COMPLETED")
    print("=" * 60)

    print(f"Total resumes: {len(resume_files)}")
    print(f"Successful: {len(candidate_results)}")
    print(f"Failed: {len(failed_resumes)}")

    print("\nCANDIDATE RANKING")
    print("-" * 60)

    if candidate_results:
        for candidate in candidate_results:
            name = (
                candidate["candidate"]["name"]
                or candidate["file_name"]
            )

            score = candidate["matching_result"]["score"]

            print(
                f"Rank {candidate['rank']}: "
                f"{name} - {score}/100"
            )
    else:
        print("No candidates were processed successfully.")

    print("\nOutput files:")
    print(f"Job details: {JOB_JSON_PATH}")
    print(f"Candidate results: {RESULTS_JSON_PATH}")

    print("\nNote: Scores are preliminary screening indicators.")
    print("Review the original resumes before making decisions.")


# ============================================================
# 11. ENTRY POINT
# ============================================================

if __name__ == "__main__":
    main()
